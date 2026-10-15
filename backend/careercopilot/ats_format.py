"""ATS format audit and ATS-friendly export (single-column DOCX / PDF)."""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import textutil as tu
from .model import ResumeModel

STANDARD_HEADINGS = {"summary": "Summary", "experience": "Experience", "projects": "Projects", "skills": "Skills",
                     "education": "Education", "certifications": "Certifications", "other": "Additional Information"}
SINGLE_DATE = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}"
PRIVATE_USE = re.compile(r"[-\U0001F300-\U0001FAFF☀-➿]")


def _f(severity: str, title: str, detail: str, fix: str = "", ok: bool = False) -> dict:
    return {"severity": "pass" if ok else severity, "title": title, "detail": detail, "fix": "" if ok else fix}


def audit(model: ResumeModel, fmt: str, layout=None) -> dict:
    """Return format findings. 'pass' items are kept so the user sees what is already right."""
    F = []
    text = model.full_text()
    F.append(_f("error", "Readable text", "The text could be extracted directly from the file." if not model.scanned else
                f"This is a scanned image; ATS software must use OCR (confidence {model.ocr_confidence or 0:.0f}%).",
                "Export the resume from Word/Google Docs as a text-based PDF or DOCX.", ok=not model.scanned))
    cols = (model.meta or {}).get("columns", 1)
    F.append(_f("warning", "Single-column layout", "Single column: reading order is unambiguous." if cols <= 1 else
                "Two or more columns were detected. Many ATS read left-to-right across columns and mix sections.",
                "Use the ATS-friendly export below, or move side-column content into the main column.", ok=cols <= 1))
    tables = (model.meta or {}).get("tables", 0)
    F.append(_f("warning", "No layout tables", "No tables are used for layout." if not tables else
                f"{tables} table(s) found. Some ATS skip or scramble text inside tables.",
                "Replace tables with plain paragraphs (the ATS-friendly export does this).", ok=not tables))
    header_contact = any(b.id.startswith(("header", "footer")) for b in model.blocks if b.section == "contact")
    F.append(_f("warning", "Contact details in the body", "Contact details are in the main body." if not header_contact else
                "Contact details are in the page header/footer, which some ATS ignore.",
                "Move your name, email and phone into the first lines of the document.", ok=not header_contact))
    has_email, has_phone = bool(tu.EMAIL_RE.search(text)), bool(tu.PHONE_RE.search(text))
    F.append(_f("error", "Email and phone found", "Email and phone number were detected." if has_email and has_phone else
                "Missing: " + ", ".join(x for x, ok in (("email", has_email), ("phone number", has_phone)) if not ok) + ".",
                "Add them as plain text near the top.", ok=has_email and has_phone))
    sections = {b.section for b in model.blocks if b.kind == "heading"}
    missing = [STANDARD_HEADINGS[s] for s in ("experience", "education", "skills") if s not in sections]
    F.append(_f("warning", "Standard section headings", "Experience, Education and Skills headings were recognised."
                if not missing else f"Not recognised: {', '.join(missing)}.",
                "Use standard headings such as 'Experience', 'Education', 'Skills', 'Projects'.", ok=not missing))
    styles = set()
    for e in model.entries.values():
        for b in model.entry_blocks(e.id):
            for m in tu.DATE_RANGE_RE.finditer(b.text):
                s = m.group("start")
                styles.add("month-name" if re.match(r"[A-Za-z]", s) else "numeric" if "/" in s else "year-only")
    F.append(_f("info", "Consistent date format", "Dates use one consistent format." if len(styles) <= 1 else
                f"Mixed date styles ({', '.join(sorted(styles))}).", "Use one style, e.g. 'Jan 2024 – May 2024'.",
                ok=len(styles) <= 1))
    long_paras = [b for b in model.blocks if b.section in ("experience", "projects") and b.kind == "paragraph"
                  and len(b.text.split()) > 40]
    F.append(_f("info", "Achievements as bullet points", "Experience and projects use short lines or bullets."
                if not long_paras else f"{len(long_paras)} long paragraph(s) in Experience/Projects.",
                "Split long paragraphs into 2–4 bullets starting with action verbs; recruiters and ATS scan bullets faster.",
                ok=not long_paras))
    pages = model.page_count or 1
    F.append(_f("info", "Length", f"{pages} page(s)." if pages <= 2 else f"{pages} pages is long for most roles.",
                "Keep to 1 page (students/entry level) or 2 pages (experienced).", ok=pages <= 2))
    odd = sorted({c for c in text if PRIVATE_USE.match(c)})
    F.append(_f("warning", "No icons or special symbols", "No icon fonts or emoji found." if not odd else
                f"{len(odd)} icon/symbol character(s) found; ATS may read them as garbage.",
                "Replace icons with words (e.g. 'Email:', 'Phone:').", ok=not odd))
    errors = sum(f["severity"] == "error" for f in F)
    warnings = sum(f["severity"] == "warning" for f in F)
    return {"findings": F, "ats_friendly": errors == 0 and warnings == 0, "errors": errors, "warnings": warnings}


# ---------------------------------------------------------------------- export
def _bottom_border(paragraph, color="888888"):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", color)):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    ppr.append(bdr)


def export_docx(draft: dict, font: str = "Calibri", size: float = 10, margin: float = 0.7) -> bytes:
    """Single-column, table-free, header-free DOCX that ATS parsers read reliably."""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(margin)
    sec.top_margin = sec.bottom_margin = Inches(max(0.4, margin - 0.1))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = font, Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.space_after = Pt(1)
    st.paragraph_format.line_spacing = 1.0
    from docx.enum.text import WD_TAB_ALIGNMENT
    right_edge = sec.page_width - sec.left_margin - sec.right_margin
    for s in draft["sections"]:
        lines = [l for l in s["lines"] if l["text"].strip()]
        if s["type"] == "contact":
            for i, l in enumerate(lines):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(l["text"])
                if i == 0:
                    r.bold, r.font.size = True, Pt(size + 6)
            continue
        if not lines:
            continue
        h = doc.add_paragraph()
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(6), Pt(2)
        r = h.add_run((s["title"] or STANDARD_HEADINGS.get(s["type"], s["type"].title())).upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size + 1), RGBColor(0x22, 0x22, 0x22)
        _bottom_border(h)
        for l in lines:
            if l["kind"] == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(l["text"])
            elif l["kind"] == "entry_header":
                p = doc.add_paragraph()
                m = re.search(r"\s*(" + tu.DATE_RANGE_RE.pattern + r"|" + SINGLE_DATE + r")\s*$", l["text"], re.I)
                if m:
                    # title on the left, dates right-aligned with a tab stop (plain text, ATS-safe)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.tab_stops.add_tab_stop(right_edge, WD_TAB_ALIGNMENT.RIGHT)
                    p.add_run(l["text"][:m.start()].strip()).bold = True
                    p.add_run("\t" + m.group(1).strip())
                else:
                    p.add_run(l["text"])
            elif l["kind"] == "skills_line" and ":" in l["text"]:
                label, rest = l["text"].split(":", 1)
                p = doc.add_paragraph()
                p.add_run(label.strip() + ": ").bold = True
                p.add_run(rest.strip())
            else:
                doc.add_paragraph(l["text"])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _soffice() -> str | None:
    for c in (os.environ.get("SOFFICE_PATH"), shutil.which("soffice"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and os.path.exists(c):
            return c
    return None


def docx_to_pdf(docx: bytes) -> bytes:
    exe = _soffice()
    if not exe:
        raise RuntimeError("PDF export needs LibreOffice; download the DOCX instead.")
    d = tempfile.mkdtemp(prefix="cc-")
    try:
        src = os.path.join(d, "resume.docx")
        open(src, "wb").write(docx)
        subprocess.run([exe, f"-env:UserInstallation=file://{d}/profile", "--headless", "--norestore",
                        "--convert-to", "pdf", "--outdir", d, src], check=True, capture_output=True, timeout=90,
                       env={**os.environ, "HOME": d})
        return open(os.path.join(d, "resume.pdf"), "rb").read()
    finally:
        shutil.rmtree(d, ignore_errors=True)


def pdf_pages(pdf: bytes) -> int:
    import pymupdf as fitz
    with fitz.open(stream=pdf, filetype="pdf") as d:
        return len(d)


def export_fitted(draft: dict, target_pages: int) -> tuple[bytes, bytes | None]:
    """Export DOCX (+PDF when LibreOffice is available), shrinking slightly until it fits `target_pages`."""
    last = None
    for size, margin in ((10, 0.7), (9.5, 0.6), (9, 0.5)):
        docx = export_docx(draft, size=size, margin=margin)
        try:
            pdf = docx_to_pdf(docx)
        except Exception:
            return docx, None
        last = (docx, pdf)
        if pdf_pages(pdf) <= target_pages:
            break
    return last
