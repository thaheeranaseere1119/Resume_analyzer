"""Cover letters: draft (Claude or template), fact-check, edit, export.

The same evidence rules as resume edits apply: every skill and number in the
letter must be supported by the resume, the candidate's notes, or an explicit
confirmation. Unsupported sentences are flagged and block export until edited.
"""
from __future__ import annotations

import io
import re
import secrets
import time

from docx import Document
from docx.shared import Pt

from . import llm
from . import textutil as tu
from .factual import _numbers
from .grammar import validate_text
from .scoring import EVIDENCE_LEVELS, collect_evidence
from .taxonomy import taxonomy

TONES = ("professional", "warm", "confident", "concise")


def candidate_name(model) -> str:
    for b in model.blocks:
        if b.section == "contact" and 1 <= len(b.text.split()) <= 4 and not re.search(r"[@\d|]", b.text):
            return b.text.title() if b.text.isupper() else b.text
    return ""


def contact_line(model) -> str:
    return next((b.text for b in model.blocks if b.section == "contact" and tu.EMAIL_RE.search(b.text)), "")


def resume_text(model) -> str:
    return "\n".join(f"[{b.section}] {b.text}" for b in model.blocks if b.kind != "contact")


def _template(model, jd, score, company, role, tone, hiring_manager, name) -> dict:
    ev = score.evidence
    req_ids = [r.skill_id for r in jd.skills if r.importance == "required"]
    strong = [s for s in req_ids if s in ev and ev[s].level in ("demonstrated", "contextual")][:3]
    tax = taxonomy()
    role = role or jd.title or "this role"
    company_txt = company or "your team"
    opener = {
        "warm": f"I was excited to see the {role} opening at {company_txt}, and I would love to bring my experience to the team.",
        "confident": f"I am applying for the {role} position at {company_txt}, where my background maps directly to what you need.",
        "concise": f"I am applying for the {role} position at {company_txt}.",
    }.get(tone, f"I am writing to apply for the {role} position at {company_txt}.")
    summary = next((b.text.rstrip(".") for b in model.blocks if b.section == "summary" and b.kind == "paragraph"), "")
    if summary and not re.match(r"(I|My)\b", summary):
        art = "an" if summary[:1].lower() in "aeiou" else "a"
        opener += f" I am {art} {summary[0].lower() + summary[1:]}."
    paras = [opener]
    evidence = []
    seen_quotes = set()
    for sid in strong:
        quote = ev[sid].quote.rstrip(".")
        verb = tu.first_word(quote)
        if not quote or quote in seen_quotes or tu.verb_tense(verb) != "past":
            continue
        seen_quotes.add(quote)
        evidence.append(f"Your posting emphasizes {tax.get(sid).name}, and I {quote[0].lower() + quote[1:]}.")
    if evidence:
        paras.append(" ".join(evidence))
    else:
        paras.append("My experience building and shipping production software has prepared me to contribute quickly.")
    paras.append(f"I would welcome the chance to discuss how I can contribute to {company_txt}. Thank you for your time and consideration.")
    return {"greeting": f"Dear {hiring_manager or 'Hiring Manager'},", "paragraphs": paras,
            "closing": f"Sincerely,\n{name}".strip()}


def _sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]


def validate_letter(letter: dict, model, confirmations: dict, notes: str) -> dict:
    """Flag unsupported skills/metrics and writing errors, sentence by sentence."""
    tax = taxonomy()
    ev = collect_evidence(model)
    known_nums = _numbers(model.full_text()) | _numbers(notes) | _numbers(" ".join(c.get("text", "") for c in confirmations.values()))
    note_skills = tax.ids_in(notes) | {s for c in confirmations.values() for s in tax.ids_in(c.get("text", ""))} | set(confirmations)
    issues = []
    for pi, para in enumerate(letter["paragraphs"]):
        for sent in _sentences(para):
            for sid in tax.ids_in(sent):
                e = ev.get(sid)
                if sid in note_skills or (e and EVIDENCE_LEVELS[e.level] > 0):
                    continue
                issues.append({"paragraph": pi, "sentence": sent, "type": "unsupported_skill", "severity": "fail",
                               "message": f"{tax.get(sid).name} is not supported by your resume or notes."})
            for n in _numbers(sent) - known_nums:
                if re.fullmatch(r"(19|20)\d{2}", n):
                    continue
                issues.append({"paragraph": pi, "sentence": sent, "type": "unsupported_metric", "severity": "fail",
                               "message": f"The figure {n} does not appear in your resume or notes."})
        g = validate_text(para, kind="letter", known_words=set(re.findall(r"[A-Za-z][A-Za-z'-]+", model.full_text() + " " + notes)))
        for c in g["checks"]:
            if c["status"] == "fail" and c["name"] not in ("keyword_stuffing",):
                issues.append({"paragraph": pi, "sentence": "", "type": c["name"], "severity": "fail", "message": c["message"]})
            elif c["status"] == "warn" and c["name"] in ("conciseness", "redundant_wording"):
                issues.append({"paragraph": pi, "sentence": "", "type": c["name"], "severity": "warn", "message": c["message"]})
    words = sum(len(p.split()) for p in letter["paragraphs"])
    return {"passed": not any(i["severity"] == "fail" for i in issues), "issues": issues, "word_count": words}


def generate(session, *, tone: str = "professional", length: str = "standard", hiring_manager: str = "",
             notes: str = "") -> dict:
    _, model, _, score = session.current()
    jd = session.jd
    company, role = session.state.get("company", ""), session.state.get("role", "") or jd.title
    name = candidate_name(model)
    writer, writer_note = "template", None
    letter = None
    if llm.configured():
        try:
            letter = llm.write_cover_letter(resume_text=resume_text(model), jd_text=session.state["jd_text"],
                                            company=company, role=role, tone=tone, length=length,
                                            hiring_manager=hiring_manager, notes=notes, candidate_name=name)
            writer = "ai"
        except llm.LLMUnavailable as exc:
            writer_note = str(exc)
    else:
        writer_note = "No AI writer is configured (add GROQ_API_KEY to .env); used the built-in writer."
    if letter is None:
        letter = _template(model, jd, score, company, role, tone, hiring_manager, name)
    letter = {"id": secrets.token_hex(4), "created": time.time(), "updated": time.time(), "tone": tone,
              "length": length, "hiring_manager": hiring_manager, "notes": notes, "writer": writer,
              "writer_label": llm.label() if writer == "ai" else "",
              "writer_note": writer_note, "resume_version": session.store.head,
              "greeting": llm.clean_text(letter["greeting"]),
              "paragraphs": [llm.clean_text(p) for p in letter["paragraphs"] if p.strip()],
              "closing": "\n".join(llm.clean_text(x) for x in letter["closing"].splitlines())}
    letter["validation"] = validate_letter(letter, model, session.state["confirmations"], notes)
    return letter


def export_docx(letter: dict, model, source_docx: bytes | None) -> bytes:
    """Plain, ATS-friendly letter. Uses the resume's own body font when the resume is a DOCX."""
    font, size = "Calibri", 11.0
    if source_docx:
        try:
            normal = Document(io.BytesIO(source_docx)).styles["Normal"].font
            font = normal.name or font
            size = normal.size.pt if normal.size else size
        except Exception:
            pass
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = font, Pt(size)
    st.paragraph_format.space_after = Pt(10)
    name = candidate_name(model)
    if name:
        r = doc.add_paragraph().add_run(name)
        r.bold, r.font.size = True, Pt(size + 4)
    contact = contact_line(model)
    if contact:
        doc.add_paragraph(contact)
    doc.add_paragraph(time.strftime("%B %d, %Y"))
    doc.add_paragraph(letter["greeting"])
    for p in letter["paragraphs"]:
        doc.add_paragraph(p)
    for line in letter["closing"].split("\n"):
        doc.add_paragraph(line).paragraph_format.space_after = Pt(0)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_text(letter: dict) -> str:
    return "\n\n".join([letter["greeting"], *letter["paragraphs"], letter["closing"]]) + "\n"
