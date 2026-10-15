"""Security: hostile content inside uploaded documents must never become active HTML or instructions."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from careercopilot.preview import docx_preview


def _hostile_docx() -> bytes:
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("hello")
    r.font.name = 'Arial" onmouseover="alert(1)'
    rid = p.part.relate_to("javascript:alert(document.cookie)", RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    rr = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "<img src=x onerror=alert(2)>"
    rr.append(t)
    h.append(rr)
    p._p.append(h)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_docx_preview_neutralises_xss_vectors():
    out = docx_preview(_hostile_docx())
    assert "javascript:" not in out
    assert "<img" not in out and "&lt;img" in out          # text is escaped, not rendered
    # no attribute injection: every attribute value stays inside its quotes
    assert not re.search(r'\s(on\w+)\s*=\s*["\']', out)
    assert "<a " not in out                                  # unsafe link dropped entirely


def test_safe_links_are_kept_with_noopener():
    d = Document()
    p = d.add_paragraph()
    rid = p.part.relate_to("https://github.com/example", RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    rr = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "GitHub"
    rr.append(t)
    h.append(rr)
    p._p.append(h)
    b = io.BytesIO()
    d.save(b)
    out = docx_preview(b.getvalue())
    assert 'href="https://github.com/example"' in out and "noopener" in out
