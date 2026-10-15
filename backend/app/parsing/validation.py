"""Untrusted upload validation. Nothing in an uploaded file is executed."""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from ..core.config import get_settings
from ..core.errors import AppError

ALLOWED = {
    "pdf": {"application/pdf", "application/x-pdf", "application/octet-stream"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream",
             "application/zip"},
}
MAX_ZIP_ENTRIES = 2000
MAX_UNCOMPRESSED = 80 * 1024 * 1024  # zip-bomb guard


@dataclass
class ValidatedUpload:
    data: bytes
    file_format: str          # pdf | docx
    mime_type: str
    filename: str
    page_count: int | None
    needs_ocr: bool
    text_chars: int


def _fail(code: str, msg: str, status: int = 400) -> AppError:
    return AppError(status, code, msg)


def validate_upload(data: bytes, filename: str, declared_mime: str | None) -> ValidatedUpload:
    s = get_settings()
    name = (filename or "resume").replace("\\", "/").split("/")[-1][:200]
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in ALLOWED:
        raise _fail("unsupported_format", "Only PDF and DOCX resumes are supported.", 415)
    if not data:
        raise _fail("empty_file", "The uploaded file is empty.")
    if len(data) > s.max_upload_bytes:
        raise _fail("file_too_large", f"The file is larger than {s.max_upload_mb} MB.", 413)
    mime = (declared_mime or "application/octet-stream").split(";")[0].strip().lower()
    if mime not in ALLOWED[ext]:
        raise _fail("mime_mismatch", "The file type does not match its extension.", 415)
    return _validate_pdf(data, name) if ext == "pdf" else _validate_docx(data, name)


def _validate_pdf(data: bytes, name: str) -> ValidatedUpload:
    if not data[:1024].lstrip().startswith(b"%PDF-"):
        raise _fail("invalid_signature", "This file is not a valid PDF.", 415)
    import pymupdf as fitz
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise _fail("corrupt_document", "The PDF could not be opened. It may be damaged.")
    try:
        if doc.needs_pass or doc.is_encrypted:
            raise _fail("encrypted_document", "The PDF is password-protected. Remove the password and upload again.")
        if len(doc) == 0:
            raise _fail("empty_document", "The PDF has no pages.")
        if len(doc) > 20:
            raise _fail("too_many_pages", "Resumes longer than 20 pages are not supported.")
        chars = sum(len(p.get_text().strip()) for p in doc)
        images = sum(len(p.get_images()) for p in doc)
        needs_ocr = chars < 40 * len(doc) and images > 0
        if chars == 0 and images == 0:
            raise _fail("empty_document", "The PDF contains no text or images.")
        return ValidatedUpload(data, "pdf", "application/pdf", name, len(doc), needs_ocr, chars)
    finally:
        doc.close()


def _validate_docx(data: bytes, name: str) -> ValidatedUpload:
    if not data.startswith(b"PK\x03\x04"):
        raise _fail("invalid_signature", "This file is not a valid DOCX document.", 415)
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        infos = z.infolist()
    except zipfile.BadZipFile:
        raise _fail("corrupt_document", "The DOCX could not be opened. It may be damaged.")
    names = {i.filename for i in infos}
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
        raise _fail("corrupt_document", "The DOCX structure is invalid.")
    if "[Content_Types].xml" not in names or "word/document.xml" not in names:
        raise _fail("invalid_signature", "This file is not a Word document.", 415)
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise _fail("macros_not_allowed", "Documents containing macros are not accepted. Save as a regular .docx.", 415)
    if b"EncryptedPackage" in data[:4096]:
        raise _fail("encrypted_document", "The document is password-protected.")
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
    except Exception:
        raise _fail("corrupt_document", "The DOCX could not be read.")
    text = "\n".join(p.text for p in doc.paragraphs) + "".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    if not text.strip():
        raise _fail("empty_document", "The document contains no text.")
    return ValidatedUpload(data, "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           name, None, False, len(text))
