"""Regression tests: only necessary keywords, correct placement, edits never blocked by pre-existing issues,
same-font PDF edits."""
import io

import pytest
from docx import Document

from careercopilot import llm
from careercopilot.service import Session

from conftest import fixture_bytes

JD = """Software Engineer

Requirements:
- Strong programming skills and good communication skills.
- Experience with Docker and software testing.
"""


@pytest.fixture(autouse=True)
def no_claude(monkeypatch):
    monkeypatch.setattr(llm, "configured", lambda: False)


def _docx(lines) -> bytes:
    d = Document()
    for style, text in lines:
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        p.add_run(text).bold = style == "Heading 1"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


RESUME = _docx([
    (None, "Asha Rao"), (None, "asha@example.com | 555-010-2030"),
    ("Heading 1", "EXPERIENCE"), (None, "Developer Intern, Acme Labs Jan 2024 - Jun 2024"),
    ("List Bullet", "Built REST services in Python and Java for the inventory team."),
    ("Heading 1", "PROJECTS"), (None, "Hotel Booking App Jan 2023 - May 2023"),
    ("List Bullet", "Developed a hotel booking application with room search and billing."),
    ("Heading 1", "SKILLS"),
    (None, "Design Tools: Figma, Canva"),
    (None, "Technical skills: Python ,Java ,SQL ,Javascript"),
    ("Heading 1", "EDUCATION"), (None, "B.Tech in Information Technology, 2020 - 2024"),
])


def _session(data_dir):
    return Session.create(RESUME, "asha.docx", JD, "Acme", "Software Engineer")


def test_generic_keywords_are_not_recommended(data_dir):
    s = _session(data_dir)
    req = {r["name"]: r["status"] for r in s.summary()["score"]["requirements"]}
    assert req["Programming"] == "met"          # shown through Python/Java, no keyword needed
    keywords = [r["keyword"] for r in s.recommendations()]
    assert "Programming" not in keywords and "Communication skills" not in keywords
    assert "Docker" in keywords and "Software testing" in keywords


def test_answer_is_placed_under_the_project_it_describes(data_dir):
    s = _session(data_dir)
    r = next(x for x in s.recommendations() if x["keyword"] == "Software testing")
    r = s.answer(r["id"], True, "I wrote test cases for the billing module of the hotel booking app and fixed the bugs I found")
    opt = r["options"][0]
    assert "Hotel Booking App" in opt["target_label"]
    assert opt["proposed_text"].startswith("Wrote test cases") and " I " not in opt["proposed_text"]
    assert opt["can_apply"]
    res = s.apply(r["id"], None, approved=True, version=s.store.head)
    assert res["comparison"]["match"] and res["preservation"]["exact"]


def test_skill_goes_to_technical_line_and_old_typos_do_not_block(data_dir):
    s = _session(data_dir)
    d = next(x for x in s.recommendations() if x["keyword"] == "Docker")
    d = s.answer(d["id"], True, "I containerized the inventory services with Docker")
    line = next(o for o in d["options"] if o["strategy"] == "skills_line")
    assert line["proposed_text"].startswith("Technical skills:") and line["proposed_text"].endswith("Docker")
    # "Python ,Java" spacing and "Javascript" were already in the user's line: reported, not blocking
    assert line["can_apply"]
    assert any("Already in your original line" in c["message"] for c in line["grammar"]["checks"])


def test_pdf_subset_font_edits_use_the_same_font():
    from datetime import date

    from careercopilot.edit_pdf import apply_pdf_edit, verify_pdf_edit
    from careercopilot.pdf_io import read_pdf
    data = fixture_bytes("two_column.pdf")  # ArialMT embedded as a subset
    m, lay = read_pdf(data, date(2026, 10, 15))
    b = next(x for x in m.blocks if "Containerized" in x.text)
    op = {"mode": "replace", "block_id": b.id, "text": "Containerized three legacy services with Docker and migrated them to the cloud."}
    out, geom = apply_pdf_edit(data, lay, op)
    assert geom["fonts"]["ArialMT"]["status"] == "exact"
    assert verify_pdf_edit(data, out, lay, op, geom)["exact"]
