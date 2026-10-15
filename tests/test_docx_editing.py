"""DOCX: formatting preservation, style inheritance, correct placement, validity."""
import io
import zipfile

import pytest
from docx import Document
from lxml import etree

from careercopilot.docx_io import W_P, W_R, body_paragraphs, para_text, qn, read_docx
from careercopilot.edit_docx import DocxEditError, apply_docx_edit, verify_docx_edit

from conftest import AS_OF, fixture_bytes

DOCX = ["classic.docx", "two_column_table.docx", "section_columns.docx"]


def _model(data):
    return read_docx(data, AS_OF)


def _bullet(model, contains):
    return next(b for b in model.blocks if b.kind == "bullet" and contains in b.text)


def _paras(data):
    return body_paragraphs(Document(io.BytesIO(data)))


@pytest.mark.parametrize("name", DOCX)
def test_replace_preserves_everything_else(name):
    data = fixture_bytes(name)
    m = _model(data)
    b = _bullet(m, "Containerized")
    new = "Containerized three legacy services with Docker and migrated them to the cloud."
    op = {"mode": "replace", "block_id": b.id, "text": new, "expected_text": b.body}
    out, info = apply_docx_edit(data, op)
    report = verify_docx_edit(data, out, op, info)
    assert report["exact"], [p for p in report["properties"] if p["status"] != "preserved"]
    # Every package part except document.xml is byte-identical.
    with zipfile.ZipFile(io.BytesIO(data)) as a, zipfile.ZipFile(io.BytesIO(out)) as z:
        assert [i.filename for i in a.infolist()] == [i.filename for i in z.infolist()]
        for n in a.namelist():
            if n != "word/document.xml":
                assert a.read(n) == z.read(n), n
    # Manual bullet glyph prefix (two_column_table uses "–\t") is kept.
    m2 = _model(out)
    edited = next(x for x in m2.blocks if x.id == b.id)
    assert edited.body == new
    assert edited.bullet_glyph == b.bullet_glyph


@pytest.mark.parametrize("name", DOCX)
def test_insert_inherits_neighbour_style(name):
    data = fixture_bytes(name)
    m = _model(data)
    anchor = _bullet(m, "Containerized")
    text = "Deployed containerized services with Docker to AWS."
    op = {"mode": "insert_after", "block_id": anchor.id, "text": text}
    out, info = apply_docx_edit(data, op)
    report = verify_docx_edit(data, out, op, info)
    assert report["exact"], [p for p in report["properties"] if p["status"] != "preserved"]
    before, after = _paras(data), _paras(out)
    idx = int(anchor.id[1:])
    tmpl, new = before[idx], after[idx + 1]
    c14n = lambda e: etree.tostring(e, method="c14n") if e is not None else b""
    # Paragraph properties (style, numbering/bullets, indentation, spacing, tabs) inherited.
    assert c14n(tmpl.find(qn("w:pPr"))) == c14n(new.find(qn("w:pPr")))
    # Body run formatting inherited from the neighbour's body text run.
    tmpl_rprs = {c14n(r.find(qn("w:rPr"))) for r in tmpl.iter(W_R)}
    new_rprs = {c14n(r.find(qn("w:rPr"))) for r in new.iter(W_R)}
    assert new_rprs <= tmpl_rprs
    # Inserted into the selected section, directly after the anchor, in the same entry.
    m2 = _model(out)
    ids = [x.id for x in m2.blocks]
    new_block = m2.blocks[ids.index(f"p{idx + 1}")]
    assert new_block.body == text and new_block.kind == "bullet"
    assert new_block.section == anchor.section and new_block.entry_id == anchor.entry_id
    assert len(after) == len(before) + 1


def test_mixed_run_formatting_kept_on_replace():
    """'40%' is bold in classic.docx; editing elsewhere in the sentence keeps it bold."""
    data = fixture_bytes("classic.docx")
    m = _model(data)
    b = _bullet(m, "PostgreSQL schemas")
    new = b.body.replace("PostgreSQL schemas", "PostgreSQL (SQL) schemas")
    out, info = apply_docx_edit(data, {"mode": "replace", "block_id": b.id, "text": new, "expected_text": b.body})
    p = _paras(out)[int(b.id[1:])]
    bold = [r for r in p.iter(W_R) if r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:b")) is not None]
    assert ["".join(t.text for t in r.iter(qn("w:t"))) for r in bold] == ["40%"]
    assert para_text(p) == new


def test_tables_hyperlinks_headers_page_breaks_preserved():
    data = fixture_bytes("classic.docx")
    m = _model(data)
    skills = next(b for b in m.blocks if b.kind == "skills_line" and "PostgreSQL" in b.text)
    assert skills.in_table
    new = skills.body + ", SQL"
    out, info = apply_docx_edit(data, {"mode": "replace", "block_id": skills.id, "text": new, "expected_text": skills.body})
    rep = {p["property"]: p for p in verify_docx_edit(data, out, {"mode": "replace", "text": new}, info)["properties"]}
    for prop in ("Tables", "Header and footer content", "Hyperlink targets", "Explicit page breaks",
                 "Page dimensions, margins and column layout"):
        assert rep[prop]["status"] == "preserved", prop
    d = Document(io.BytesIO(out))
    assert d.tables[0].rows[2].cells[1].text == "Git, PostgreSQL, Linux, SQL"
    assert "github.com/jordanlee" in "".join(p.text for p in d.sections[0].header.paragraphs) or \
        b"github.com/jordanlee" in zipfile.ZipFile(io.BytesIO(out)).read("word/header1.xml")


def test_section_columns_unchanged():
    data = fixture_bytes("section_columns.docx")
    m = _model(data)
    assert m.meta["columns"] == 2
    line = next(b for b in m.blocks if b.kind == "skills_line" and "Tools" in b.text)
    out, info = apply_docx_edit(data, {"mode": "replace", "block_id": line.id, "text": line.body + ", Docker",
                                       "expected_text": line.body})
    assert read_docx(out, AS_OF).meta["columns"] == 2
    a = [etree.tostring(s, method="c14n") for s in Document(io.BytesIO(data)).element.body.iter(qn("w:sectPr"))]
    b = [etree.tostring(s, method="c14n") for s in Document(io.BytesIO(out)).element.body.iter(qn("w:sectPr"))]
    assert a == b


def test_stale_target_refused():
    data = fixture_bytes("classic.docx")
    m = _model(data)
    b = _bullet(m, "Containerized")
    with pytest.raises(DocxEditError):
        apply_docx_edit(data, {"mode": "replace", "block_id": b.id, "text": "x", "expected_text": "something else"})


def test_insert_never_duplicates_section_break():
    data = fixture_bytes("section_columns.docx")
    root = Document(io.BytesIO(data)).element.body
    with_sect = [i for i, p in enumerate(root.iter(W_P)) if p.find(qn("w:pPr")) is not None
                 and p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
    assert with_sect
    i = with_sect[0]
    out, _ = apply_docx_edit(data, {"mode": "insert_after", "block_id": f"p{i}", "text": "Inserted line."})
    count = lambda d: len(list(Document(io.BytesIO(d)).element.body.iter(qn("w:sectPr"))))
    assert count(out) == count(data)
