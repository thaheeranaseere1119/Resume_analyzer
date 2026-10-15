"""Wording/grammar scan, ATS format audit, editor (edit + reorder + write back), ATS export, eligible jobs."""
import io
from datetime import date

import pymupdf as fitz
import pytest
from docx import Document

from careercopilot import llm
from careercopilot.polish import improve_line
from careercopilot.service import Session

from conftest import fixture_bytes


@pytest.fixture(autouse=True)
def no_claude(monkeypatch):
    monkeypatch.setattr(llm, "configured", lambda: False)


@pytest.mark.parametrize("before,kind,after", [
    ("Technical skills: Python ,Java ,Javascript ,MongoDb,JQuery,Flutter(basics)", "skills_line",
     "Technical skills: Python, Java, JavaScript, MongoDB, jQuery, Flutter (basics)"),
    ("Responsible for managing the college website.", "bullet", "Managed the college website."),
    ("Structured backend workflows and Designing a user-friendly interface.", "paragraph",
     "Structured backend workflows and designed a user-friendly interface."),
    ("Outcome:presented research paper at the conference.", "paragraph", "Outcome: Presented research paper at the conference."),
    ("Built lots of features with docker.", "bullet", "Built many features with Docker."),
    ("Python with Data Science by csc", "entry_header", "Python with Data Science by csc"),
    ("Urdu (Native) • Tamil (Native)", "paragraph", "Urdu (Native) • Tamil (Native)"),
])
def test_professional_wording(before, kind, after):
    assert improve_line(before, kind=kind, tense="past") == after


def _session(name="single_column.pdf"):
    return Session.create(fixture_bytes(name), name, (fixture_bytes("job_description.txt")).decode(), "Acme", "Backend")


def test_wording_cards_apply_with_exact_formatting(data_dir):
    d = Document()
    for style, text in [(None, "Asha Rao"), (None, "asha@example.com | 555-010-2030"), ("Heading 1", "EXPERIENCE"),
                        (None, "Intern, Acme Labs Jan 2024 - Jun 2024"),
                        ("List Bullet", "Responsible for managing the inventory dashboard used by the sales team."),
                        ("Heading 1", "SKILLS"), (None, "Technical skills: Python ,Java ,Javascript ,MongoDb"),
                        ("Heading 1", "EDUCATION"), (None, "B.Tech in Information Technology 2020 - 2024")]:
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        p.add_run(text)
    buf = io.BytesIO()
    d.save(buf)
    s = Session.create(buf.getvalue(), "asha.docx", (fixture_bytes("job_description.txt")).decode(), "Acme", "Backend")
    w = {x["options"][0]["original_text"]: x for x in s.recommendations() if x["kind"] == "wording"}
    skills = w["Technical skills: Python ,Java ,Javascript ,MongoDb"]["options"][0]
    assert skills["proposed_text"] == "Technical skills: Python, Java, JavaScript, MongoDB"
    bullet = w["Responsible for managing the inventory dashboard used by the sales team."]["options"][0]
    assert bullet["proposed_text"] == "Managed the inventory dashboard used by the sales team."
    assert skills["factual"]["status"] == bullet["factual"]["status"] == "stylistic"
    rid = w["Technical skills: Python ,Java ,Javascript ,MongoDb"]["id"]
    res = s.apply(rid, None, approved=True, version=s.store.head)
    assert res["preservation"]["exact"]


def test_format_audit_flags_two_columns_and_tables(data_dir):
    a = _session("two_column.pdf").format_audit()
    assert any(f["title"] == "Single-column layout" and f["severity"] == "warning" for f in a["findings"])
    b = _session("two_column_table.docx").format_audit()
    assert any(f["title"] == "No layout tables" and f["severity"] == "warning" for f in b["findings"])
    c = _session("single_column.pdf").format_audit()
    assert c["ats_friendly"]


def test_editor_reorder_and_ats_export(data_dir):
    s = _session("classic.docx")
    p = s.editor_get()
    d = p["draft"]
    types = [x["type"] for x in d["sections"]]
    # drag Skills above Experience and move a bullet between entries
    i_sk, i_ex = types.index("skills"), types.index("experience")
    d["sections"].insert(i_ex, d["sections"].pop(i_sk))
    exp = next(x for x in d["sections"] if x["type"] == "experience")
    exp["lines"].insert(1, exp["lines"].pop())
    p = s.editor_save(d)
    assert [x["type"] for x in p["draft"]["sections"]].index("skills") < [x["type"] for x in p["draft"]["sections"]].index("experience")
    docx, name, _ = s.editor_export("docx")
    doc = Document(io.BytesIO(docx))
    text = "\n".join(par.text for par in doc.paragraphs)
    assert text.index("SKILLS") < text.index("EXPERIENCE") and not doc.tables and name.endswith(".docx")
    pdf, _, _ = s.editor_export("pdf")
    with fitz.open(stream=pdf, filetype="pdf") as f:
        assert "Containerized three legacy services" in "".join(pg.get_text() for pg in f)


def test_editor_improve_and_write_back(data_dir):
    s = _session("classic.docx")
    r = s.editor_improve("i was responsible for managing our docker builds", "bullet")
    assert r["changed"] and r["text"][0].isupper() and " i " not in f" {r['text']} "
    d = s.editor_get()["draft"]
    line = next(l for sec in d["sections"] for l in sec["lines"] if "Wrote automated tests" in l["text"])
    line["text"] = "Wrote automated tests in Python that raised coverage from 45% to 80%."
    s.editor_save(d)
    out = s.editor_apply_to_original()
    assert out["applied"] and out["head"] == 1
    blocks = [b.body for b in s.current()[1].blocks]
    assert "Wrote automated tests in Python that raised coverage from 45% to 80%." in blocks


def test_eligible_jobs(data_dir):
    jobs = _session("classic.docx").eligible_jobs()
    titles = {j["title"]: j for j in jobs}
    assert titles["Python Developer"]["eligible"]
    assert not titles["UI/UX Designer"]["eligible"]
    assert all("linkedin.com" in j["search"]["LinkedIn"] for j in jobs)
    assert jobs == sorted(jobs, key=lambda j: (not j["eligible"], -j["score"]))


def test_ats_explanation_is_consistent_with_score(data_dir):
    s = _session()
    d = s.ats_explained()
    comps = [c for c in d["score"]["components"].values() if c["applicable"]]
    tw = sum(c["weight"] for c in comps)
    assert round(sum(100 * c["score"] * c["weight"] / tw for c in comps), 1) == d["score"]["overall"]
    assert d["skills"] and d["format"]["findings"]


def test_negative_answers_never_become_claims(data_dir):
    from careercopilot.phrasing import phrase_bullets
    s = _session("classic.docx")
    r = next(x for x in s.recommendations() if x["keyword"] == "AWS")
    r = s.answer(r["id"], True, "I set up Linux virtual machines in a lab. While I have not deployed anything on AWS, "
                                "I am eager to learn AWS services.")
    assert r["status"] == "unsupported" and r["options"] == []
    lines = phrase_bullets("I configured Linux VMs for a lab. I have not used AWS yet.", skill="AWS", tense="past")
    assert lines and all("AWS" not in l for l in lines)


def test_unnamed_font_spans_do_not_fail_verification():
    from careercopilot.edit_pdf import apply_pdf_edit, verify_pdf_edit
    from careercopilot.pdf_io import read_pdf
    import pymupdf as fitz
    # build a page whose text uses a font object without a usable name, then edit another line
    base = fitz.open(stream=fixture_bytes("single_column.pdf"), filetype="pdf")
    data = base.tobytes()
    m, lay = read_pdf(data, date(2026, 10, 15))
    b = next(x for x in m.blocks if "Containerized" in x.text)
    op = {"mode": "replace", "block_id": b.id, "text": "Containerized three legacy services with Docker."}
    out1, g1 = apply_pdf_edit(data, lay, op)
    m2, lay2 = read_pdf(out1, date(2026, 10, 15))
    for sp, _ in [(s, 0) for blk in lay2.blocks.values() for l in blk.lines for s in l.spans]:
        if "Docker" in sp["text"]:
            sp["font"] = ""  # simulate a font reported without a name
    b2 = next(x for x in m2.blocks if "Wrote automated tests" in x.text)
    op2 = {"mode": "replace", "block_id": b2.id, "text": "Wrote automated tests that raised coverage from 45% to 80% in Python."}
    out2, g2 = apply_pdf_edit(out1, lay2, op2)
    rep = verify_pdf_edit(out1, out2, lay2, op2, g2)
    assert next(p for p in rep["properties"] if p["property"].startswith("Unaffected"))["status"] == "preserved"


def test_inserting_a_copy_of_an_existing_line_is_blocked(data_dir):
    s = _session("single_column.pdf")
    rec = s.recommender()
    b = next(x for x in s.current()[1].blocks if "Containerized" in x.text)
    opt = rec.evaluate({"mode": "insert_after", "action": "add", "block_id": b.id, "proposed_text": b.body,
                        "strategy": "editor"})
    assert not opt["can_apply"] and "already in your resume" in opt["blocked_reason"]


def test_wrapped_bullet_lines_are_joined():
    from careercopilot.pdf_io import read_pdf
    import pymupdf as fitz
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((50, 60), "EXPERIENCE", fontname="hebo", fontsize=11)
    p.insert_text((50, 80), "PR Head                                                  Aug 2025 - Present", fontname="hebo", fontsize=9)
    p.insert_text((50, 96), "•", fontsize=9)
    p.insert_text((60, 96), "Led outreach for departmental events, coordinating student teams to improve event", fontsize=9)
    p.insert_text((50, 108), "participation by 30%.", fontsize=9)   # wrapped line starting at the bullet's x
    m, _ = read_pdf(doc.tobytes(), date(2026, 10, 15))
    texts = [b.text for b in m.blocks]
    assert any(t.endswith("improve event participation by 30%.") for t in texts)
    assert not any(t.startswith("participation") for t in texts)


def test_wording_fix_raises_the_score(data_dir):
    d = Document()
    for style, text in [(None, "Asha Rao"), (None, "asha@example.com | 555-010-2030"), ("Heading 1", "EXPERIENCE"),
                        (None, "Intern, Acme Labs Jan 2024 - Jun 2024"),
                        ("List Bullet", "Responsible for managing the inventory dashboard used by the sales team."),
                        ("Heading 1", "SKILLS"), (None, "Technical skills: Python ,Java ,Javascript ,MongoDb"),
                        ("Heading 1", "EDUCATION"), (None, "B.Tech in Information Technology 2020 - 2024")]:
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        p.add_run(text)
    buf = io.BytesIO()
    d.save(buf)
    s = Session.create(buf.getvalue(), "asha.docx", (fixture_bytes("job_description.txt")).decode(), "Acme", "Backend")
    w = next(x for x in s.recommendations() if x["kind"] == "wording")
    pts = w["options"][0]["score"]["points"]
    assert pts > 0
    res = s.apply(w["id"], None, approved=True, version=s.store.head)
    assert res["comparison"]["actual_points"] == pts and res["comparison"]["match"]


def test_any_one_of_skill_lists_are_alternatives(jd_text, weights):
    from careercopilot.docx_io import read_docx
    from careercopilot.jd import parse_jd
    from careercopilot.scoring import score_resume
    from conftest import AS_OF
    jd = parse_jd("Requirements:\n"
                  "- Build Python services and APIs using FastAPI, Flask, or Django.\n"
                  "- Vector databases (FAISS, Pinecone, Weaviate, Azure AI Search, pgvector, or similar).\n"
                  "- Deploy on Azure ML, AWS SageMaker, and GCP Vertex AI.\n"
                  "- Version control (Git) and Docker.")
    by = {r.skill_id: r for r in jd.skills}
    assert by["flask"].example_only and set(by["flask"].alternatives) == {"fastapi", "django"}
    assert by["aws"].example_only and set(by["aws"].alternatives) == {"azure", "gcp"}
    assert not by["git"].example_only and not by["docker"].example_only and not by["python"].example_only
    # A resume with Flask satisfies FastAPI and Django too.
    m = read_docx(fixture_bytes("classic.docx"), AS_OF)
    b = next(x for x in m.blocks if x.kind == "bullet")
    m2 = m.with_replaced(b.id, "Built a Flask service that exposed inventory data to the sales team.")
    reqs = {r["skill_id"]: r for r in score_resume(m2, jd, weights).requirements if r["type"] == "skill"}
    assert reqs["flask"]["status"] == "met"
    assert reqs["django"]["status"] == "met" and reqs["django"]["via"] == "flask"
