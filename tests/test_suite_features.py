"""Tracker, cover letters and the Claude writing layer (Claude is mocked; no network)."""
import io
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from careercopilot import llm, tracker
from careercopilot.api import app
from careercopilot.grammar import validate_text
from careercopilot.service import Session, WorkflowError

from conftest import fixture_bytes


@pytest.fixture(autouse=True)
def no_real_claude(monkeypatch):
    monkeypatch.setattr(llm, "configured", lambda: False)
    llm._cache.clear()


def _session(jd_text, name="classic.docx"):
    return Session.create(fixture_bytes(name), name, jd_text, "Northwind", "Senior Backend Engineer")


# ------------------------------------------------------------------ tracker
def test_tracker_crud_and_history(data_dir):
    j = tracker.create_job({"company": "Acme", "title": "Engineer"})
    assert j["status"] == "saved" and len(j["history"]) == 1
    j = tracker.update_job(j["id"], {"status": "applied", "notes": "Sent via referral"})
    assert j["applied_on"] and [h["status"] for h in j["history"]] == ["saved", "applied"]
    with pytest.raises(tracker.TrackerError):
        tracker.update_job(j["id"], {"status": "ghosted"})
    with pytest.raises(tracker.TrackerError):
        tracker.create_job({"notes": "nothing else"})
    assert tracker.delete_job(j["id"]) == {"deleted": j["id"]}
    assert tracker.list_jobs() == []


# ------------------------------------------------------------- cover letters
def test_template_letter_is_supported_and_exportable(data_dir, jd_text):
    s = _session(jd_text)
    letter = s.create_letter(tone="professional", length="standard", hiring_manager="Priya Shah", notes="")
    assert letter["writer"] == "template" and "GROQ_API_KEY" in letter["writer_note"]
    assert letter["greeting"] == "Dear Priya Shah,"
    assert letter["validation"]["passed"], letter["validation"]
    data, name, _ = s.export_letter(letter["id"], "docx")
    doc = Document(io.BytesIO(data))
    assert name.endswith(".docx") and doc.styles["Normal"].font.name == "Calibri"  # resume's own body font
    assert "Jordan Lee" in doc.paragraphs[0].text
    txt, _, _ = s.export_letter(letter["id"], "txt")
    assert txt.decode().startswith("Dear Priya Shah,")


def test_letter_edits_with_unsupported_claims_block_export(data_dir, jd_text):
    s = _session(jd_text)
    letter = s.create_letter()
    original = list(letter["paragraphs"])
    paras = original + ["I ran Kubernetes clusters that cut costs by 63%."]
    upd = s.update_letter(letter["id"], letter["greeting"], paras, letter["closing"])
    types = {i["type"] for i in upd["validation"]["issues"]}
    assert {"unsupported_skill", "unsupported_metric"} <= types and not upd["validation"]["passed"]
    with pytest.raises(WorkflowError) as e:
        s.export_letter(letter["id"], "docx")
    assert e.value.status == 409
    # A fact the candidate supplied in their own notes is allowed.
    s.state["letters"][0]["notes"] = "I mentored two interns and ran Kubernetes clusters at a hackathon."
    upd = s.update_letter(letter["id"], letter["greeting"], original + ["I have also run Kubernetes clusters."], letter["closing"])
    assert upd["validation"]["passed"]


def test_letters_allow_first_person_but_bullets_do_not():
    assert validate_text("I built order services in Python for the team.", kind="letter")["passed"]
    assert not validate_text("I built order services in Python for the team.", kind="bullet")["passed"]


# ------------------------------------------------------------ Claude layer
class FakeClient:
    def __init__(self, text, stop_reason="end_turn"):
        self.calls = []
        msg = SimpleNamespace(stop_reason=stop_reason, content=[SimpleNamespace(type="text", text=text)])
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: (self.calls.append(kw), msg)[1]))


def test_generate_json_request_shape_and_refusal(monkeypatch):
    monkeypatch.setattr(llm, "provider", lambda: "anthropic")
    fake = FakeClient('{"text": "Built a thing."}')
    monkeypatch.setattr(llm, "_get_client", lambda: fake)
    out = llm.generate_json("prompt", llm.BULLET_SCHEMA)
    assert out == {"text": "Built a thing."}
    kw = fake.calls[0]
    assert kw["model"] == "claude-opus-5-5" and kw["fallbacks"] == "default"
    assert kw["betas"] == ["server-side-fallback-2026-07-01"]
    assert kw["output_config"]["format"]["type"] == "json_schema"
    llm._cache.clear()
    monkeypatch.setattr(llm, "_get_client", lambda: FakeClient("", stop_reason="refusal"))
    with pytest.raises(llm.LLMUnavailable):
        llm.generate_json("prompt", llm.BULLET_SCHEMA)


def test_claude_draft_kept_only_when_it_passes_validation(data_dir, jd_text, monkeypatch):
    s = _session(jd_text)
    j = next(r for r in s.recommendations() if r["keyword"] == "Java")
    monkeypatch.setattr(llm, "configured", lambda: True)
    good = "Built a Java Spring Boot service that syncs inventory with the warehouse API."
    monkeypatch.setattr(llm, "write_bullet", lambda **kw: good)
    s.answer(j["id"], True, "i built a java spring boot service for inventory sync with the warehouse API")
    opt = next(r for r in s.recommendations() if r["keyword"] == "Java")["options"][0]
    assert opt["writer"] == "ai" and opt["proposed_text"] == good and opt["can_apply"]
    # A hallucinated draft (new tool + invented metric) is rejected; the template draft stays.
    monkeypatch.setattr(llm, "write_bullet",
                        lambda **kw: "Built a Java Spring Boot service on Kubernetes, cutting latency by 73%.")
    opt = next(r for r in s.recommendations() if r["keyword"] == "Java")["options"][0]
    assert opt["writer"] == "phraser" and "did not pass" in opt["writer_note"] and opt["can_apply"]
    assert opt["proposed_text"] == "Built a Java Spring Boot service for inventory sync with the warehouse API."


def test_claude_letter_and_outage_fallback(data_dir, jd_text, monkeypatch):
    s = _session(jd_text)
    monkeypatch.setattr(llm, "configured", lambda: True)
    monkeypatch.setattr(llm, "write_cover_letter", lambda **kw: {
        "greeting": "Dear Hiring Manager,", "closing": "Sincerely,\nJordan Lee",
        "paragraphs": ["I am excited to apply for the Senior Backend Engineer role at Northwind.",
                       "At Acme Commerce I built order-processing services in Python and Flask that handle 2M requests per day."]})
    l = s.create_letter()
    assert l["writer"] == "ai" and l["validation"]["passed"]

    def down(**kw):
        raise llm.LLMUnavailable("Could not reach the Claude API")
    monkeypatch.setattr(llm, "write_cover_letter", down)
    l2 = s.create_letter()
    assert l2["writer"] == "template" and "Could not reach" in l2["writer_note"]


def test_claude_rewrite_requires_configuration(data_dir, jd_text):
    s = _session(jd_text)
    sql = next(r for r in s.recommendations() if r["keyword"] == "SQL")
    with pytest.raises(WorkflowError) as e:
        s.claude_rewrite(sql["id"], None)
    assert e.value.status == 409


# ---------------------------------------------------------------------- API
def test_http_tracker_letters_and_session_list(data_dir, jd_text):
    c = TestClient(app)
    r = c.post("/api/sessions", files={"file": ("classic.docx", fixture_bytes("classic.docx"))},
               data={"jd_text": jd_text, "company": "Northwind", "role": "Backend Engineer", "track": "true"})
    sid = r.json()["session_id"]
    jobs = c.get("/api/jobs").json()
    assert len(jobs) == 1 and jobs[0]["session"]["session_id"] == sid and jobs[0]["company"] == "Northwind"
    assert c.patch(f"/api/jobs/{jobs[0]['id']}", json={"status": "interview"}).json()["status"] == "interview"
    assert c.patch(f"/api/jobs/{jobs[0]['id']}", json={"status": "nope"}).status_code == 400
    assert [s["session_id"] for s in c.get("/api/sessions").json()] == [sid]
    letter = c.post(f"/api/sessions/{sid}/letters", json={"tone": "warm"}).json()
    assert c.get(f"/api/sessions/{sid}/letters/{letter['id']}/download?format=txt").status_code == 200
    assert c.post(f"/api/sessions/{sid}/letters", json={"tone": "sarcastic"}).status_code == 400
    assert c.get("/api/meta").json()["statuses"] == tracker.STATUSES
    assert c.delete(f"/api/jobs/{jobs[0]['id']}").status_code == 200


# ---------------------------------------------------------------------- Groq
class _Resp:
    def __init__(self, payload):
        self.payload = payload

    def read(self):
        import json as _j
        return _j.dumps(self.payload).encode()

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _groq_reply(content, finish="stop"):
    return _Resp({"choices": [{"finish_reason": finish, "message": {"content": content}}]})


def test_groq_is_used_when_key_is_set(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "gsk_test")
    sent = []
    monkeypatch.setattr(llm.urllib.request, "urlopen",
                        lambda req, timeout: (sent.append(req), _groq_reply('{"text": "Built a thing."}'))[1])
    assert llm.provider() == "groq" and llm.label() == "Groq" and llm.model_name() == llm.GROQ_MODEL
    assert llm.generate_json("prompt", llm.BULLET_SCHEMA) == {"text": "Built a thing."}
    import json as _j
    body = _j.loads(sent[0].data)
    assert sent[0].full_url == llm.GROQ_URL and body["response_format"] == {"type": "json_object"}
    assert sent[0].get_header("Authorization") == "Bearer gsk_test"
    llm._cache.clear()
    monkeypatch.setattr(llm.urllib.request, "urlopen", lambda req, timeout: _groq_reply("not json"))
    with pytest.raises(llm.LLMUnavailable):
        llm.generate_json("prompt", llm.BULLET_SCHEMA)


def test_invented_details_are_rejected():
    src = "i was responsible for managing our docker builds and fixed lots of bugs in it"
    bad = "Managed Docker builds and resolved numerous bugs, ensuring stable and efficient container deployment."
    assert llm.grounding_problems(bad, [src])
    assert not llm.grounding_problems("Managed Docker builds and fixed numerous bugs.", [src])
    notes = "i dockerized 3 of our old services with docker compose so they run the same on every laptop and in CI"
    assert not llm.grounding_problems("Dockerized three legacy services using Docker Compose, enabling identical "
                                      "operation on laptops and CI pipelines.",
                                      [notes, "Containerized three legacy services and migrated them to the cloud."])
    assert not llm.grounding_problems("Improved website speed and reduced bugs.",
                                      ["Worked on making the website more fast and less buggy"], max_new=1)


def test_ai_wording_cards_only_when_score_rises(data_dir, jd_text, monkeypatch):
    import io
    from docx import Document
    d = Document()
    for style, text in [(None, "Asha Rao"), (None, "asha@example.com | 555-010-2030"), ("Heading 1", "EXPERIENCE"),
                        (None, "Intern, Acme Labs Jan 2024 - Jun 2024"),
                        ("List Bullet", "Built a API that handle orders for the the sales team using Python."),
                        ("List Bullet", "Wrote automated tests that raised coverage from 45% to 80%."),
                        ("Heading 1", "SKILLS"), (None, "Python, SQL, Docker"),
                        ("Heading 1", "EDUCATION"), (None, "B.Tech in Information Technology 2020 - 2024")]:
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        p.add_run(text)
    buf = io.BytesIO()
    d.save(buf)
    s = Session.create(buf.getvalue(), "asha.docx", jd_text)
    blocks = {b.body: b.id for b in s.current()[1].blocks}
    bad, ok = blocks["Built a API that handle orders for the the sales team using Python."], \
        blocks["Wrote automated tests that raised coverage from 45% to 80%."]
    monkeypatch.setattr(llm, "configured", lambda: True)
    monkeypatch.setattr(llm, "polish_lines", lambda lines: {
        bad: "Built an API that handles orders for the sales team using Python.",
        ok: "Wrote automated tests that increased coverage from 45% to 80%."})   # correct already: no gain
    recs = {r["options"][0]["block_id"]: r["options"][0] for r in s.recommendations() if r["kind"] == "wording"}
    fixed = recs[bad]
    assert fixed["writer"] == "ai" and fixed["proposed_text"] == "Built an API that handles orders for the sales team using Python."
    assert fixed["score"]["points"] > 0 and fixed["can_apply"]
    assert ok not in recs                      # a rewording that cannot raise the score is not offered
    assert all(o["score"]["points"] > 0 for o in recs.values())
